"""Data loaders for analysis inputs.

Discovers paths under collected-data/; does not hardcode pilot US lists.
Falls back to the consolidated PRR workbook when annotation CSVs are empty shells.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd

from analysis import config


def _require_file(path: Path, label: str) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"Required input missing ({label}): {path}")
    return path


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def find_prr_workbook(data_dir: Path | None = None) -> Path | None:
    """Return the newest matching consolidated PRR workbook, if any."""
    data_dir = data_dir or config.DATA_DIR
    if not data_dir.is_dir():
        return None
    matches = sorted(data_dir.glob(config.PRR_WORKBOOK_GLOB), key=lambda p: p.stat().st_mtime)
    return matches[-1] if matches else None


def _sheet_to_frame(workbook_path: Path, sheet_name: str) -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise FileNotFoundError(f"Sheet {sheet_name!r} not found in {workbook_path}")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()
    header = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records: list[dict[str, str]] = []
    for raw in rows[1:]:
        if raw is None or raw[0] is None or str(raw[0]).strip() == "":
            continue
        rec = {}
        for i, col in enumerate(header):
            val = raw[i] if i < len(raw) else None
            rec[col] = "" if val is None else str(val).strip()
        records.append(rec)
    return pd.DataFrame(records)


def _mapped_fill_count(df: pd.DataFrame, col: str = "mapped_gap_ids") -> int:
    if df.empty or col not in df.columns:
        return 0
    return int(df[col].astype(str).str.strip().ne("").sum())


def load_annotations_from_workbook(workbook_path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Load Evaluator-1, Evaluator-2 and Adjudicação sheets.

    Adjudicação is normalized to disagreements columns expected by mappings.py
    (blind_item_id, final_gap_ids, …). Does not alter the workbook.
    """
    e1 = _sheet_to_frame(workbook_path, "Evaluator-1")
    e2 = _sheet_to_frame(workbook_path, "Evaluator-2")
    adj = _sheet_to_frame(workbook_path, "Adjudicação")

    for label, df in (("Evaluator-1", e1), ("Evaluator-2", e2)):
        if "mapped_gap_ids" not in df.columns:
            raise ValueError(f"{label} missing mapped_gap_ids in {workbook_path}")

    # Normalize adjudication → disagreements schema
    if not adj.empty:
        rename = {}
        if "final_gap_ids" not in adj.columns:
            raise ValueError(f"Adjudicação missing final_gap_ids in {workbook_path}")
        keep = [
            c
            for c in [
                "blind_item_id",
                "user_story_id",
                "question_text_raw",
                "evaluator_1_gap_ids",
                "evaluator_2_gap_ids",
                "final_gap_ids",
                "notes",
                "justificativa_adjudicacao",
                "status",
            ]
            if c in adj.columns
        ]
        adj = adj[keep].copy()
        if "notes" not in adj.columns and "justificativa_adjudicacao" in adj.columns:
            adj["notes"] = adj["justificativa_adjudicacao"]
    else:
        adj = pd.DataFrame(
            columns=[
                "blind_item_id",
                "user_story_id",
                "question_text_raw",
                "evaluator_1_gap_ids",
                "evaluator_2_gap_ids",
                "final_gap_ids",
                "notes",
            ]
        )

    return e1, e2, adj


@dataclass(frozen=True)
class AnalysisInputs:
    questions: pd.DataFrame
    evaluator_1: pd.DataFrame
    evaluator_2: pd.DataFrame
    disagreements: pd.DataFrame
    blind_id_map: pd.DataFrame
    prr_reference: pd.DataFrame
    prr_gap_states: pd.DataFrame
    run_summary: pd.DataFrame
    paths: dict[str, Path]
    annotation_source: str = "csv"

    def input_hashes(self) -> dict[str, str]:
        return {name: file_sha256(path) for name, path in self.paths.items() if path.is_file()}


def load_analysis_inputs(
    *,
    collected: Path | None = None,
) -> AnalysisInputs:
    """Load final collected-data artifacts used by the analysis pipeline."""
    collected = collected or config.COLLECTED
    annotation = collected / "annotation"
    private = annotation / "private"
    reference = collected / "reference"
    audit = collected / "audit"

    paths: dict[str, Path] = {
        "questions": _require_file(collected / "questions.csv", "questions"),
        "evaluator_1": _require_file(annotation / "evaluator-1.csv", "evaluator-1"),
        "evaluator_2": _require_file(annotation / "evaluator-2.csv", "evaluator-2"),
        "disagreements": _require_file(annotation / "disagreements.csv", "disagreements"),
        "blind_id_map": _require_file(private / "blind-id-map.csv", "blind-id-map"),
        "prr_reference": _require_file(reference / "prr-reference.csv", "prr-reference"),
        "prr_gap_states": _require_file(reference / "prr-gap-states.csv", "prr-gap-states"),
        "run_summary": _require_file(audit / "run-summary.csv", "run-summary"),
    }

    e1 = read_csv(paths["evaluator_1"])
    e2 = read_csv(paths["evaluator_2"])
    disag = read_csv(paths["disagreements"])
    sources: list[str] = []

    csv_e_filled = min(_mapped_fill_count(e1), _mapped_fill_count(e2))
    csv_adj_filled = _mapped_fill_count(disag, col="final_gap_ids")
    if csv_e_filled > 0:
        sources.append("csv:evaluators")
    if csv_adj_filled > 0:
        sources.append("csv:disagreements")

    workbook = find_prr_workbook()
    if workbook is not None:
        paths["prr_workbook"] = workbook
        need_evaluators = csv_e_filled == 0
        need_adjudication = csv_adj_filled == 0
        if need_evaluators or need_adjudication:
            x_e1, x_e2, x_adj = load_annotations_from_workbook(workbook)
            if need_evaluators and _mapped_fill_count(x_e1) > 0 and _mapped_fill_count(x_e2) > 0:
                e1, e2 = x_e1, x_e2
                sources.append(f"xlsx:evaluators:{workbook.name}")
            if need_adjudication and _mapped_fill_count(x_adj, col="final_gap_ids") > 0:
                disag = x_adj
                sources.append(f"xlsx:adjudicacao:{workbook.name}")

    annotation_source = "+".join(sources) if sources else "csv:empty"

    return AnalysisInputs(
        questions=read_csv(paths["questions"]),
        evaluator_1=e1,
        evaluator_2=e2,
        disagreements=disag,
        blind_id_map=read_csv(paths["blind_id_map"]),
        prr_reference=read_csv(paths["prr_reference"]),
        prr_gap_states=read_csv(paths["prr_gap_states"]),
        run_summary=read_csv(paths["run_summary"]),
        paths=paths,
        annotation_source=annotation_source,
    )


def environment_versions() -> dict[str, Any]:
    import sys

    import matplotlib
    import numpy
    import pandas

    versions: dict[str, Any] = {
        "python": sys.version.split()[0],
        "pandas": pandas.__version__,
        "numpy": numpy.__version__,
        "matplotlib": matplotlib.__version__,
    }
    try:
        import scipy

        versions["scipy"] = scipy.__version__
    except ImportError:
        versions["scipy"] = "not installed"
    try:
        import statsmodels

        versions["statsmodels"] = statsmodels.__version__
    except ImportError:
        versions["statsmodels"] = "not installed"
    return versions
