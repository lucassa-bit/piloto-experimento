#!/usr/bin/env python3
"""
Gera materials/, manifest.csv, runs.csv e validation_report.md a partir do PRR.

Fonte de verdade: data/Matrizes de rastreabilidade - PRR.xlsx
Não altera a planilha. Não inventa, resume ou parafraseia informações de referência.
Não executa Spec Kit / Codex / runner.
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EXPERIMENT_DIR = Path(__file__).resolve().parent
DEFAULT_DATA = EXPERIMENT_DIR / "data" / "Matrizes de rastreabilidade - PRR.xlsx"
DEFAULT_SOURCE_COPY_FROM = Path(
    "/home/lucassa-bit/Downloads/Matrizes de rastreabilidade - PRR (1).xlsx"
)

PRR_SHEETS = ("PRR - US02", "PRR - US08", "PRR - US18", "PRR - US25")
USER_STORY_IDS = ("US02", "US08", "US18", "US25")
CONDITIONS = ("C0", "CL", "CO", "CD", "CS", "CT")
CATEGORY_TO_CONDITION = {
    "Léxica": "CL",
    "Operacional": "CO",
    "Decisória": "CD",
    "Sistêmica": "CS",
}
CONDITION_TO_CATEGORY = {v: k for k, v in CATEGORY_TO_CONDITION.items()}
# Ordem dos blocos em total.md (espelha experimento-extendido)
CATEGORY_BLOCK_ORDER = ("Léxica", "Operacional", "Decisória", "Sistêmica")
CATEGORY_FILE = {
    "Léxica": "lexical.md",
    "Operacional": "operational.md",
    "Decisória": "decisional.md",
    "Sistêmica": "systemic.md",
}
CATEGORY_TITLE_PT = {
    "Léxica": "Contexto léxico",
    "Operacional": "Contexto operacional",
    "Decisória": "Contexto decisório",
    "Sistêmica": "Contexto sistêmico",
}
CATEGORY_HEADING_EN = {
    "Léxica": "Lexical Context",
    "Operacional": "Operational Context",
    "Decisória": "Decisional Context",
    "Sistêmica": "Systemic Context",
}
CONDITION_CONTEXT_FILES = {
    "CL": "lexical.md",
    "CO": "operational.md",
    "CD": "decisional.md",
    "CS": "systemic.md",
    "CT": "total.md",
}
MATRIX_COLS = {"C0": 11, "CL": 12, "CO": 13, "CD": 14, "CS": 15, "CT": 16}
REPETITIONS = 3

STATE_OPEN = "Aberta"
STATE_ANSWERED = "Respondida"
GAP_STATE_EN = {STATE_OPEN: "Open", STATE_ANSWERED: "Answered"}

MATERIAL_FILES = (
    "user-story.md",
    "user-story-original.md",
    "lexical.md",
    "operational.md",
    "decisional.md",
    "systemic.md",
    "total.md",
)


@dataclass
class GapRow:
    gap_id: str
    lacuna: str | None
    importance: str | None
    ref_id: str | None
    reference_information: str | None
    category: str | None
    justificativa: str | None
    gap_id_matrix: str | None
    category_matrix: str | None
    states: dict[str, str | None]


@dataclass
class UserStoryPRR:
    user_story_id: str
    user_story_text: str | None
    source: str | None = None
    gaps: list[GapRow] = field(default_factory=list)


@dataclass
class ValidationIssue:
    user_story_id: str
    gap_id: str
    condition: str
    problem: str


def _as_str(value: Any) -> str | None:
    if value is None:
        return None
    return str(value)


def ensure_data_file(data_path: Path, source_path: Path | None) -> None:
    """Garante cópia local da planilha sem modificar o arquivo de origem."""
    data_path.parent.mkdir(parents=True, exist_ok=True)
    if data_path.exists():
        return
    if source_path is None or not source_path.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada em {data_path} e fonte {source_path} indisponível."
        )
    shutil.copy2(source_path, data_path)


def load_user_stories_catalog(workbook) -> dict[str, str]:
    if "user_stories" not in workbook.sheetnames:
        return {}
    ws = workbook["user_stories"]
    catalog: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        us_id = ws.cell(row, 1).value
        text = ws.cell(row, 3).value
        if us_id and text:
            catalog[str(us_id).strip()] = str(text)
    return catalog


def load_materials_sources(workbook) -> dict[str, str]:
    """Lê Source da aba materials (sem inventar Dataset/Arquivo/Linha)."""
    if "materials" not in workbook.sheetnames:
        return {}
    ws = workbook["materials"]
    sources: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        us_id = ws.cell(row, 1).value
        source = ws.cell(row, 2).value
        if not us_id or not source:
            continue
        key = str(us_id).strip()
        # Primeira ocorrência prevalece (há IDs duplicados na aba)
        if key not in sources:
            sources[key] = str(source).strip()
    return sources


def parse_prr_sheet(workbook, sheet_name: str, sources: dict[str, str]) -> UserStoryPRR:
    ws = workbook[sheet_name]
    us_id = sheet_name.split(" - ", 1)[1].strip()
    gaps: list[GapRow] = []
    user_story_text: str | None = None

    for row in range(2, ws.max_row + 1):
        label = ws.cell(row, 2).value
        if isinstance(label, str) and label.strip().startswith("Original User Story"):
            nxt = ws.cell(row + 1, 2).value
            if nxt:
                user_story_text = str(nxt)
            continue

        gap_id = ws.cell(row, 1).value
        if gap_id is None:
            continue
        gap_id_str = str(gap_id).strip()
        if not re.match(r"^G\d+$", gap_id_str):
            continue

        states = {
            cond: (
                None
                if ws.cell(row, col).value is None
                else str(ws.cell(row, col).value).strip()
            )
            for cond, col in MATRIX_COLS.items()
        }

        gaps.append(
            GapRow(
                gap_id=gap_id_str,
                lacuna=_as_str(ws.cell(row, 2).value),
                importance=_as_str(ws.cell(row, 3).value),
                ref_id=_as_str(ws.cell(row, 4).value),
                reference_information=_as_str(ws.cell(row, 5).value),
                category=_as_str(ws.cell(row, 6).value),
                justificativa=_as_str(ws.cell(row, 7).value),
                gap_id_matrix=_as_str(ws.cell(row, 9).value),
                category_matrix=_as_str(ws.cell(row, 10).value),
                states=states,
            )
        )

    return UserStoryPRR(
        user_story_id=us_id,
        user_story_text=user_story_text,
        source=sources.get(us_id),
        gaps=gaps,
    )


def refs_for_category(prr: UserStoryPRR, category: str) -> list[GapRow]:
    """Gaps da categoria com informação de referência, na ordem do PRR."""
    return [
        g
        for g in prr.gaps
        if g.category == category
        and g.reference_information is not None
        and g.reference_information != ""
    ]


def refs_for_condition(prr: UserStoryPRR, condition: str) -> list[GapRow]:
    if condition == "C0":
        return []
    if condition == "CT":
        # União na ordem de blocos do extendido (L → O → D → S), PRR dentro de cada bloco
        result: list[GapRow] = []
        for cat in CATEGORY_BLOCK_ORDER:
            result.extend(refs_for_category(prr, cat))
        return result
    category = CONDITION_TO_CATEGORY[condition]
    return refs_for_category(prr, category)


def bullets_for_gaps(gaps: list[GapRow]) -> str:
    lines = [f"- {g.reference_information}" for g in gaps if g.reference_information]
    return "\n".join(lines)


def format_category_file(us_id: str, category: str, gaps: list[GapRow]) -> str:
    title = CATEGORY_TITLE_PT[category]
    body = bullets_for_gaps(gaps)
    if body:
        return f"# {us_id} — {title}\n\n{body}\n"
    return f"# {us_id} — {title}\n"


def format_total_file(prr: UserStoryPRR) -> str:
    parts: list[str] = []
    for cat in CATEGORY_BLOCK_ORDER:
        gaps = refs_for_category(prr, cat)
        heading = CATEGORY_HEADING_EN[cat]
        body = bullets_for_gaps(gaps)
        if body:
            parts.append(f"# {heading}\n\n{body}")
        else:
            parts.append(f"# {heading}")
    return "\n\n".join(parts) + "\n"


def format_user_story_file(
    us_id: str,
    title_suffix: str,
    user_story_text: str,
    source: str | None,
) -> str:
    lines = [f"# {us_id} — {title_suffix}", "", user_story_text]
    if source:
        lines.extend(["", "## Proveniência", "", f"- **Fonte:** {source}"])
    lines.append("")
    return "\n".join(lines)


def write_materials(prrs: list[UserStoryPRR], materials_dir: Path) -> None:
    materials_dir.mkdir(parents=True, exist_ok=True)
    for prr in prrs:
        if not prr.user_story_text:
            raise ValueError(f"{prr.user_story_id}: User Story original ausente no PRR")
        us_dir = materials_dir / prr.user_story_id
        us_dir.mkdir(parents=True, exist_ok=True)

        (us_dir / "user-story.md").write_text(
            format_user_story_file(
                prr.user_story_id, "User Story", prr.user_story_text, prr.source
            ),
            encoding="utf-8",
        )
        (us_dir / "user-story-original.md").write_text(
            format_user_story_file(
                prr.user_story_id,
                "User Story original",
                prr.user_story_text,
                prr.source,
            ),
            encoding="utf-8",
        )

        for cat in CATEGORY_BLOCK_ORDER:
            fname = CATEGORY_FILE[cat]
            content = format_category_file(
                prr.user_story_id, cat, refs_for_category(prr, cat)
            )
            (us_dir / fname).write_text(content, encoding="utf-8")

        (us_dir / "total.md").write_text(format_total_file(prr), encoding="utf-8")


def write_manifest(path: Path, prrs: list[UserStoryPRR]) -> None:
    fieldnames = [
        "user_story_id",
        "condition",
        "gap_id",
        "ref_id",
        "category",
        "importance",
        "gap_state",
        "reference_information",
    ]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for prr in prrs:
            for condition in CONDITIONS:
                for gap in prr.gaps:
                    state_pt = gap.states.get(condition)
                    writer.writerow(
                        {
                            "user_story_id": prr.user_story_id,
                            "condition": condition,
                            "gap_id": gap.gap_id,
                            "ref_id": gap.ref_id or "",
                            "category": gap.category or "",
                            "importance": gap.importance or "",
                            "gap_state": GAP_STATE_EN.get(state_pt or "", state_pt or ""),
                            "reference_information": gap.reference_information or "",
                        }
                    )


def write_runs(path: Path, prrs: list[UserStoryPRR], repetitions: int = REPETITIONS) -> None:
    fieldnames = [
        "run_id",
        "user_story_id",
        "condition",
        "repetition",
        "user_story_file",
        "context_file",
    ]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for prr in prrs:
            for condition in CONDITIONS:
                us_file = f"materials/{prr.user_story_id}/user-story.md"
                ctx_name = CONDITION_CONTEXT_FILES.get(condition)
                ctx_file = (
                    ""
                    if ctx_name is None
                    else f"materials/{prr.user_story_id}/{ctx_name}"
                )
                for rep in range(1, repetitions + 1):
                    run_id = f"{prr.user_story_id}_{condition}_R{rep}"
                    writer.writerow(
                        {
                            "run_id": run_id,
                            "user_story_id": prr.user_story_id,
                            "condition": condition,
                            "repetition": f"R{rep}",
                            "user_story_file": us_file,
                            "context_file": ctx_file,
                        }
                    )


def _parse_bullets(md_text: str) -> list[str]:
    """Extrai textos de bullets `- ...` preservando ordem."""
    bullets: list[str] = []
    for line in md_text.splitlines():
        if line.startswith("- "):
            bullets.append(line[2:])
    return bullets


def _section_bullets(md_text: str, heading: str) -> list[str]:
    """Bullets sob um heading `# Heading` até o próximo `# `."""
    pattern = re.compile(
        rf"^# {re.escape(heading)}\s*\n(.*?)(?=^# |\Z)",
        re.MULTILINE | re.DOTALL,
    )
    match = pattern.search(md_text)
    if not match:
        return []
    return _parse_bullets(match.group(1))


def validate(
    prrs: list[UserStoryPRR],
    catalog: dict[str, str],
    materials_dir: Path,
) -> tuple[list[ValidationIssue], dict[str, Any]]:
    issues: list[ValidationIssue] = []
    report: dict[str, Any] = {"per_us": {}, "global_notes": []}

    for prr in prrs:
        us = prr.user_story_id
        us_report: dict[str, Any] = {}
        us_dir = materials_dir / us

        gap_ids = [g.gap_id for g in prr.gaps]
        ref_ids = [g.ref_id for g in prr.gaps if g.ref_id]
        dup_gaps = sorted({g for g, c in Counter(gap_ids).items() if c > 1})
        dup_refs = sorted({r for r, c in Counter(ref_ids).items() if c > 1})
        us_report["duplicate_gap_ids"] = dup_gaps
        us_report["duplicate_ref_ids"] = dup_refs
        for g in dup_gaps:
            issues.append(ValidationIssue(us, g, "", "Gap ID duplicado no PRR"))
        for r in dup_refs:
            issues.append(ValidationIssue(us, "", "", f"Ref ID duplicado: {r}"))

        gaps_without_info = [g.gap_id for g in prr.gaps if not g.reference_information]
        us_report["gaps_without_reference_information"] = gaps_without_info
        for g in gaps_without_info:
            issues.append(ValidationIssue(us, g, "", "Gap sem Informação de referência"))

        for g in prr.gaps:
            if not g.ref_id:
                issues.append(ValidationIssue(us, g.gap_id, "", "Gap sem Ref ID correspondente"))
            if not g.category:
                issues.append(ValidationIssue(us, g.gap_id, "", "Gap sem categoria"))

        us_report["orphan_reference_information"] = []

        catalog_text = catalog.get(us)
        us_report["user_story_prr"] = prr.user_story_text
        us_report["user_story_catalog"] = catalog_text
        us_report["user_story_match"] = (
            catalog_text is not None and catalog_text == prr.user_story_text
        )
        if catalog_text is None:
            issues.append(
                ValidationIssue(us, "", "", "User Story ausente na aba user_stories")
            )
        elif catalog_text != prr.user_story_text:
            issues.append(
                ValidationIssue(
                    us,
                    "",
                    "",
                    "Divergência entre User Story do PRR e aba user_stories",
                )
            )

        us_report["total_gaps"] = len(prr.gaps)
        us_report["necessary"] = sum(1 for g in prr.gaps if g.importance == "Necessária")
        us_report["complementary"] = sum(
            1 for g in prr.gaps if g.importance == "Complementar"
        )
        us_report["by_category"] = dict(Counter(g.category for g in prr.gaps))

        answered: dict[str, list[str]] = {}
        open_gaps: dict[str, list[str]] = {}
        for cond in CONDITIONS:
            answered[cond] = [
                g.gap_id for g in prr.gaps if g.states.get(cond) == STATE_ANSWERED
            ]
            open_gaps[cond] = [
                g.gap_id for g in prr.gaps if g.states.get(cond) == STATE_OPEN
            ]
            for g in prr.gaps:
                st = g.states.get(cond)
                if st is None or st == "":
                    issues.append(
                        ValidationIssue(us, g.gap_id, cond, "Gap sem estado na matriz")
                    )
                elif st not in (STATE_OPEN, STATE_ANSWERED):
                    issues.append(
                        ValidationIssue(
                            us, g.gap_id, cond, f"Estado inesperado: {st!r}"
                        )
                    )

        us_report["answered_by_condition"] = answered
        us_report["open_by_condition"] = open_gaps

        cat_matrix_div: list[str] = []
        for g in prr.gaps:
            if g.category != g.category_matrix:
                msg = (
                    f"Categoria esquerda ({g.category!r}) != matriz ({g.category_matrix!r})"
                )
                cat_matrix_div.append(f"{g.gap_id}: {msg}")
                issues.append(ValidationIssue(us, g.gap_id, "", msg))
            if g.gap_id != g.gap_id_matrix:
                msg = f"Gap ID esquerdo ({g.gap_id}) != matriz ({g.gap_id_matrix})"
                cat_matrix_div.append(msg)
                issues.append(ValidationIssue(us, g.gap_id, "", msg))

            expected_cond = CATEGORY_TO_CONDITION.get(g.category or "")
            for cond in CONDITIONS:
                st = g.states.get(cond)
                if expected_cond is None:
                    continue
                if cond == "C0":
                    expected = STATE_OPEN
                elif cond == "CT":
                    expected = STATE_ANSWERED
                elif cond == expected_cond:
                    expected = STATE_ANSWERED
                else:
                    expected = STATE_OPEN
                if st != expected:
                    msg = (
                        f"Divergência categoria×matriz: esperado {expected}, "
                        f"obtido {st} (categoria={g.category})"
                    )
                    cat_matrix_div.append(f"{g.gap_id}/{cond}: {msg}")
                    issues.append(ValidationIssue(us, g.gap_id, cond, msg))

        us_report["category_matrix_divergences"] = cat_matrix_div

        # --- Materials presence ---
        for fname in MATERIAL_FILES:
            fpath = us_dir / fname
            if not fpath.is_file():
                issues.append(
                    ValidationIssue(us, "", "", f"Arquivo de material ausente: {fname}")
                )

        # C0: conceitualmente sem arquivo de contexto dedicado
        if (us_dir / "c0.md").exists() or (us_dir / "none.md").exists():
            issues.append(
                ValidationIssue(
                    us, "", "C0", "C0 não deve ter arquivo de contexto em materials/"
                )
            )

        expected_us = format_user_story_file(
            us, "User Story", prr.user_story_text or "", prr.source
        )
        expected_us_orig = format_user_story_file(
            us, "User Story original", prr.user_story_text or "", prr.source
        )
        for fname, expected in (
            ("user-story.md", expected_us),
            ("user-story-original.md", expected_us_orig),
        ):
            fpath = us_dir / fname
            if fpath.is_file() and fpath.read_text(encoding="utf-8") != expected:
                issues.append(
                    ValidationIssue(
                        us, "", "", f"Conteúdo de {fname} diverge do texto do PRR"
                    )
                )

        # Category files: only that category; exact PRR order; no Gap/Ref IDs
        category_bullets: dict[str, list[str]] = {}
        sizes: dict[str, dict[str, int]] = {}
        for cat in CATEGORY_BLOCK_ORDER:
            cond = CATEGORY_TO_CONDITION[cat]
            fname = CATEGORY_FILE[cat]
            fpath = us_dir / fname
            expected_gaps = refs_for_category(prr, cat)
            expected_texts = [
                g.reference_information for g in expected_gaps if g.reference_information
            ]
            expected_content = format_category_file(us, cat, expected_gaps)
            if fpath.is_file():
                actual = fpath.read_text(encoding="utf-8")
                if actual != expected_content:
                    issues.append(
                        ValidationIssue(
                            us, "", cond, f"Conteúdo de {fname} diverge do esperado"
                        )
                    )
                bullets = _parse_bullets(actual)
                category_bullets[cat] = bullets
                if bullets != expected_texts:
                    issues.append(
                        ValidationIssue(
                            us,
                            "",
                            cond,
                            f"Bullets de {fname} != referência da categoria na ordem do PRR",
                        )
                    )
                # Sem Gap/Ref IDs no corpo
                body = actual.split("\n", 1)[1] if "\n" in actual else ""
                for g in prr.gaps:
                    if g.gap_id and re.search(rf"\b{re.escape(g.gap_id)}\b", body):
                        issues.append(
                            ValidationIssue(
                                us, g.gap_id, cond, f"Gap ID aparece em {fname}"
                            )
                        )
                    if g.ref_id and re.search(rf"\b{re.escape(g.ref_id)}\b", body):
                        issues.append(
                            ValidationIssue(
                                us, g.gap_id, cond, f"Ref ID aparece em {fname}"
                            )
                        )
                body_join = "\n".join(bullets)
                sizes[cond] = {
                    "chars": len(body_join),
                    "words": len(body_join.split()) if body_join else 0,
                }
            else:
                category_bullets[cat] = []
                sizes[cond] = {"chars": 0, "words": 0}

        # total.md == união dos quatro (ordem de blocos extendido)
        expected_total = format_total_file(prr)
        total_path = us_dir / "total.md"
        ct_union_ok = False
        if total_path.is_file():
            actual_total = total_path.read_text(encoding="utf-8")
            if actual_total != expected_total:
                issues.append(
                    ValidationIssue(
                        us, "", "CT", "Conteúdo de total.md diverge do esperado"
                    )
                )
            # Cada seção == arquivo categórico correspondente
            section_ok = True
            for cat in CATEGORY_BLOCK_ORDER:
                heading = CATEGORY_HEADING_EN[cat]
                sec = _section_bullets(actual_total, heading)
                if sec != category_bullets.get(cat, []):
                    section_ok = False
                    issues.append(
                        ValidationIssue(
                            us,
                            "",
                            "CT",
                            f"Seção {heading} de total.md != {CATEGORY_FILE[cat]}",
                        )
                    )
            # Multiset CT == união dos quatro
            union_list: list[str] = []
            for cat in CATEGORY_BLOCK_ORDER:
                union_list.extend(category_bullets.get(cat, []))
            ct_bullets = _parse_bullets(actual_total)
            ct_union_ok = section_ok and ct_bullets == union_list
            if not ct_union_ok and section_ok:
                issues.append(
                    ValidationIssue(
                        us,
                        "",
                        "CT",
                        "CT não é exatamente a união de CL+CO+CD+CS (ordem de blocos)",
                    )
                )
            body_ct = "\n".join(ct_bullets)
            sizes["CT"] = {
                "chars": len(body_ct),
                "words": len(body_ct.split()) if body_ct else 0,
            }
            # Sem Gap/Ref IDs
            for g in prr.gaps:
                if g.gap_id and re.search(rf"\b{re.escape(g.gap_id)}\b", actual_total):
                    # só no corpo (não no heading US)
                    if re.search(
                        rf"(?m)^- .*{re.escape(g.gap_id)}", actual_total
                    ) or re.search(rf"\b{re.escape(g.gap_id)}\b", "\n".join(ct_bullets)):
                        issues.append(
                            ValidationIssue(
                                us, g.gap_id, "CT", "Gap ID aparece em total.md"
                            )
                        )
                if g.ref_id and any(g.ref_id in b for b in ct_bullets):
                    if re.search(rf"\b{re.escape(g.ref_id)}\b", "\n".join(ct_bullets)):
                        issues.append(
                            ValidationIssue(
                                us, g.gap_id, "CT", "Ref ID aparece em total.md"
                            )
                        )
        else:
            sizes["CT"] = {"chars": 0, "words": 0}

        us_report["ct_is_exact_union"] = ct_union_ok
        us_report["context_sizes"] = sizes

        # Regras 11–12: Respondida ↔ info disponível na condição (via categoria)
        for cond in CONDITIONS:
            included_gap_ids = {g.gap_id for g in refs_for_condition(prr, cond)}
            for g in prr.gaps:
                st = g.states.get(cond)
                has_info = g.gap_id in included_gap_ids
                if st == STATE_ANSWERED and not has_info and cond != "C0":
                    issues.append(
                        ValidationIssue(
                            us,
                            g.gap_id,
                            cond,
                            "Gap Respondida sem informação de referência disponível na condição",
                        )
                    )
                if st == STATE_ANSWERED and cond == "C0":
                    issues.append(
                        ValidationIssue(
                            us,
                            g.gap_id,
                            "C0",
                            "Gap marcada Respondida em C0 (C0 não possui contexto)",
                        )
                    )
                if st == STATE_OPEN and has_info:
                    issues.append(
                        ValidationIssue(
                            us,
                            g.gap_id,
                            cond,
                            "Gap Aberta possui informação de referência disponível na condição",
                        )
                    )

        report["per_us"][us] = us_report

    return issues, report


def write_validation_report(
    path: Path,
    prrs: list[UserStoryPRR],
    issues: list[ValidationIssue],
    report: dict[str, Any],
) -> None:
    lines: list[str] = []
    lines.append("# Validation Report")
    lines.append("")
    lines.append("Fonte: `data/Matrizes de rastreabilidade - PRR.xlsx`")
    lines.append("")
    lines.append("## Resumo global")
    lines.append("")
    lines.append(f"- User Stories no experimento: {len(prrs)}")
    lines.append(f"- Issues encontradas: {len(issues)}")
    lines.append("")

    if issues:
        lines.append("## Issues")
        lines.append("")
        lines.append("| user_story_id | gap_id | condition | problem |")
        lines.append("|---|---|---|---|")
        for i in issues:
            lines.append(
                f"| {i.user_story_id} | {i.gap_id} | {i.condition} | {i.problem} |"
            )
        lines.append("")
    else:
        lines.append("## Issues")
        lines.append("")
        lines.append("Nenhuma inconsistência encontrada pelas regras de validação.")
        lines.append("")

    for prr in prrs:
        us = prr.user_story_id
        r = report["per_us"][us]
        lines.append(f"## {us}")
        lines.append("")
        lines.append(f"- Quantidade total de gaps: **{r['total_gaps']}**")
        lines.append(f"- Gaps necessários (Necessária): **{r['necessary']}**")
        lines.append(f"- Gaps complementares (Complementar): **{r['complementary']}**")
        lines.append("- Informações por categoria:")
        for cat in CATEGORY_BLOCK_ORDER:
            lines.append(f"  - {cat}: {r['by_category'].get(cat, 0)}")
        lines.append("")
        lines.append("### Gaps respondidos por condição")
        lines.append("")
        for cond in CONDITIONS:
            ids = ", ".join(r["answered_by_condition"][cond]) or "(nenhum)"
            lines.append(f"- {cond}: {len(r['answered_by_condition'][cond])} — {ids}")
        lines.append("")
        lines.append("### Gaps abertos por condição")
        lines.append("")
        for cond in CONDITIONS:
            ids = ", ".join(r["open_by_condition"][cond]) or "(nenhum)"
            lines.append(f"- {cond}: {len(r['open_by_condition'][cond])} — {ids}")
        lines.append("")
        lines.append("### CT = união CL + CO + CD + CS")
        lines.append("")
        lines.append(
            f"- Confirmação: **{'SIM' if r['ct_is_exact_union'] else 'NÃO'}** "
            "(ordem de blocos Lexical→Operational→Decisional→Systemic; "
            "bullets iguais aos quatro arquivos de categoria)"
        )
        lines.append("")
        lines.append("### Comparação User Story (PRR × user_stories)")
        lines.append("")
        lines.append(f"- Match: **{'SIM' if r['user_story_match'] else 'NÃO'}**")
        lines.append(f"- PRR: {r['user_story_prr']}")
        lines.append(f"- user_stories: {r['user_story_catalog']}")
        lines.append("")
        lines.append("### Duplicatas e lacunas estruturais")
        lines.append("")
        lines.append(f"- Ref IDs duplicados: {r['duplicate_ref_ids'] or '(nenhum)'}")
        lines.append(f"- Gap IDs duplicados: {r['duplicate_gap_ids'] or '(nenhum)'}")
        lines.append(
            f"- Gaps sem informação de referência: "
            f"{r['gaps_without_reference_information'] or '(nenhum)'}"
        )
        lines.append(
            f"- Informações de referência sem gap: "
            f"{r['orphan_reference_information'] or '(nenhum)'}"
        )
        lines.append("")
        lines.append("### Divergências categoria × matriz de estados")
        lines.append("")
        if r["category_matrix_divergences"]:
            for d in r["category_matrix_divergences"]:
                lines.append(f"- {d}")
        else:
            lines.append("- Nenhuma.")
        lines.append("")
        lines.append("### Tamanho dos contextos (caracteres / palavras)")
        lines.append("")
        lines.append(
            "Diferenças de tamanho entre condições **não** são tratadas como erro; "
            "apenas registradas."
        )
        lines.append("")
        lines.append("| condition | characters | words |")
        lines.append("|---|---:|---:|")
        for cond in ("CL", "CO", "CD", "CS", "CT"):
            s = r["context_sizes"][cond]
            lines.append(f"| {cond} | {s['chars']} | {s['words']} |")
        lines.append("")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def remove_obsolete_layout(experiment_dir: Path) -> None:
    """Remove pastas USxx/Cy/input.md do layout antigo."""
    for us_id in USER_STORY_IDS:
        old = experiment_dir / us_id
        if old.is_dir():
            shutil.rmtree(old)
    obsolete = experiment_dir / "prepare_experiment.py"
    if obsolete.is_file():
        obsolete.unlink()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Gera materials/ e artefatos do experimento piloto a partir do PRR."
    )
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE_COPY_FROM)
    parser.add_argument("--repetitions", type=int, default=REPETITIONS)
    parser.add_argument("--experiment-dir", type=Path, default=EXPERIMENT_DIR)
    parser.add_argument(
        "--keep-obsolete",
        action="store_true",
        help="Não remove pastas USxx/C*/ do layout antigo",
    )
    args = parser.parse_args(argv)

    experiment_dir: Path = args.experiment_dir
    materials_dir = experiment_dir / "materials"

    ensure_data_file(args.data, args.source)

    wb = load_workbook(args.data, data_only=True, read_only=False)
    missing = [s for s in PRR_SHEETS if s not in wb.sheetnames]
    if missing:
        print(f"Abas PRR ausentes: {missing}", file=sys.stderr)
        return 1

    catalog = load_user_stories_catalog(wb)
    sources = load_materials_sources(wb)
    prrs = [parse_prr_sheet(wb, name, sources) for name in PRR_SHEETS]
    wb.close()

    write_materials(prrs, materials_dir)
    write_manifest(experiment_dir / "manifest.csv", prrs)
    write_runs(experiment_dir / "runs.csv", prrs, repetitions=args.repetitions)

    issues, report = validate(prrs, catalog, materials_dir)
    write_validation_report(
        experiment_dir / "validation_report.md", prrs, issues, report
    )

    if not args.keep_obsolete:
        remove_obsolete_layout(experiment_dir)

    n_runs = len(prrs) * len(CONDITIONS) * args.repetitions
    print(f"Materials gerados para {len(prrs)} User Stories.")
    print(f"Runs planejados: {n_runs}")
    print(f"Issues: {len(issues)}")
    print(f"Relatório: {experiment_dir / 'validation_report.md'}")
    return 0 if not issues else 2


if __name__ == "__main__":
    raise SystemExit(main())
